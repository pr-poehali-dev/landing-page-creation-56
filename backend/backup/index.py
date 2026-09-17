import json
import os
import io
import time
from datetime import datetime, date
import psycopg2
import boto3
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

STATUS_LABELS = {
    'new': 'Новая',
    'estimate': 'Смета',
    'contract': 'Договор',
    'payment': 'Оплата',
    'live': 'Размещение',
    'completed': 'Завершена',
    'lost': 'Потеряна',
}

DOC_LABELS = {
    'contract': 'Договор',
    'invoice': 'Счёт',
    'act': 'Акт',
}

SOURCE_LABELS = {
    'form': 'Форма на сайте',
    'manual': 'Добавлена вручную',
}

HEAD_FILL = PatternFill('solid', fgColor='1E293B')
HEAD_FONT = Font(bold=True, color='FFFFFF', size=11)
THIN = Side(style='thin', color='CBD5E1')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def fmt_date(v):
    if not v:
        return ''
    if isinstance(v, datetime):
        return v.strftime('%d.%m.%Y %H:%M')
    if isinstance(v, date):
        return v.strftime('%d.%m.%Y')
    return str(v)


def style_sheet(ws, widths, n_cols):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'
    for row in ws.iter_rows(min_row=2, max_col=n_cols):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical='top', wrap_text=False)


def build_workbook(leads, docs):
    wb = Workbook()

    ws = wb.active
    ws.title = 'Заявки'
    ws.append([
        'ID', 'Дата заявки', 'Статус', 'Имя', 'Телефон', 'Компания',
        'Сумма, ₽', 'Оплачено, ₽', 'Остаток, ₽',
        'Размещение, ₽', 'Ролик, ₽', 'Нужен ролик',
        'Начало', 'Окончание', 'Дней', 'Хронометраж, сек',
        'ИНН', 'КПП', 'ОГРН', 'Юр. адрес',
        'Банк', 'Расчётный счёт', 'БИК', 'Корр. счёт',
        'Подписант', 'Должность', 'Источник', 'Комментарий',
    ])

    total_sum = 0
    total_paid = 0
    for l in leads:
        price = l['total_price'] or 0
        paid = l['paid_amount'] or 0
        total_sum += price
        total_paid += paid
        ws.append([
            l['id'],
            fmt_date(l['created_at']),
            STATUS_LABELS.get(l['status'], l['status'] or ''),
            l['name'] or '',
            l['phone'] or '',
            l['company'] or '',
            price,
            paid,
            max(price - paid, 0),
            l['placement_amount'] or 0,
            l['video_amount'] or 0,
            'Да' if l['need_video'] else 'Нет',
            fmt_date(l['start_date']),
            fmt_date(l['end_date']),
            l['days'] or '',
            l['duration'] or '',
            l['inn'] or '',
            l['kpp'] or '',
            l['ogrn'] or '',
            l['legal_address'] or '',
            l['bank_name'] or '',
            l['bank_account'] or '',
            l['bank_bik'] or '',
            l['bank_corr_account'] or '',
            l['signer_name'] or '',
            l['signer_position'] or '',
            SOURCE_LABELS.get(l['source'], l['source'] or ''),
            l['comment'] or '',
        ])

    widths = [6, 17, 12, 22, 16, 26, 13, 13, 13, 15, 12, 12, 12, 12, 8, 16,
              15, 12, 16, 40, 30, 24, 12, 24, 24, 20, 18, 40]
    style_sheet(ws, widths, 28)
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=11):
        for cell in row:
            cell.number_format = '# ##0'

    ws2 = wb.create_sheet('Документы')
    ws2.append(['ID заявки', 'Клиент', 'Тип документа', 'Номер', 'Дата', 'Ссылка на файл'])
    names = {l['id']: (l['company'] or l['name'] or '') for l in leads}
    for d in docs:
        ws2.append([
            d['lead_id'],
            names.get(d['lead_id'], ''),
            DOC_LABELS.get(d['doc_type'], d['doc_type'] or ''),
            d['doc_no'] or '',
            fmt_date(d['created_at']),
            d['file_url'] or '',
        ])
    style_sheet(ws2, [11, 30, 16, 12, 17, 80], 6)
    for row in ws2.iter_rows(min_row=2, min_col=6, max_col=6):
        for cell in row:
            if cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(color='2563EB', underline='single')

    ws3 = wb.create_sheet('Сводка')
    ws3.append(['Показатель', 'Значение'])
    by_status = {}
    for l in leads:
        key = STATUS_LABELS.get(l['status'], l['status'] or '')
        by_status[key] = by_status.get(key, 0) + 1

    rows = [
        ['Дата выгрузки', datetime.now().strftime('%d.%m.%Y %H:%M')],
        ['Всего заявок', len(leads)],
        ['Всего документов', len(docs)],
        ['Общая сумма сделок, ₽', total_sum],
        ['Получено оплат, ₽', total_paid],
        ['Ожидается к оплате, ₽', max(total_sum - total_paid, 0)],
        ['', ''],
        ['Заявки по статусам', ''],
    ]
    for k, v in by_status.items():
        rows.append([k, v])
    for r in rows:
        ws3.append(r)
    style_sheet(ws3, [32, 24], 2)
    for row in ws3.iter_rows(min_row=2, max_col=1):
        for cell in row:
            cell.font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read(), total_sum


def fetch_data(dsn):
    conn = psycopg2.connect(dsn)
    cur = conn.cursor()
    cur.execute(
        "SELECT id, name, phone, comment, duration, days, need_video, total_price, source, "
        "status, created_at, company, start_date, end_date, placement_amount, video_amount, "
        "inn, kpp, ogrn, legal_address, bank_name, bank_account, bank_bik, bank_corr_account, "
        "signer_name, signer_position, paid_amount FROM leads ORDER BY id"
    )
    cols = ['id', 'name', 'phone', 'comment', 'duration', 'days', 'need_video', 'total_price',
            'source', 'status', 'created_at', 'company', 'start_date', 'end_date',
            'placement_amount', 'video_amount', 'inn', 'kpp', 'ogrn', 'legal_address',
            'bank_name', 'bank_account', 'bank_bik', 'bank_corr_account', 'signer_name',
            'signer_position', 'paid_amount']
    leads = [dict(zip(cols, r)) for r in cur.fetchall()]

    cur.execute("SELECT lead_id, doc_type, file_url, doc_no, created_at FROM lead_documents ORDER BY id")
    dcols = ['lead_id', 'doc_type', 'file_url', 'doc_no', 'created_at']
    docs = [dict(zip(dcols, r)) for r in cur.fetchall()]
    cur.close()
    conn.close()
    return leads, docs


def list_backups(dsn):
    conn = psycopg2.connect(dsn)
    cur = conn.cursor()
    cur.execute(
        "SELECT id, period, file_url, leads_count, docs_count, total_sum, created_at "
        "FROM backups ORDER BY created_at DESC LIMIT 24"
    )
    items = [{
        'id': r[0],
        'period': r[1],
        'url': r[2],
        'leadsCount': r[3],
        'docsCount': r[4],
        'totalSum': int(r[5] or 0),
        'createdAt': r[6].isoformat() if r[6] else None,
    } for r in cur.fetchall()]
    cur.close()
    conn.close()
    return items


def handler(event: dict, context) -> dict:
    """Формирует резервную выгрузку базы заявок и документов в Excel, сохраняет её в хранилище и возвращает историю архивов"""
    method = event.get('httpMethod', 'GET')

    cors_headers = {
        'Access-Control-Allow-Origin': '*',
        'Access-Control-Allow-Methods': 'GET, POST, OPTIONS',
        'Access-Control-Allow-Headers': 'Content-Type, X-Admin-Key',
        'Access-Control-Max-Age': '86400'
    }

    if method == 'OPTIONS':
        return {'statusCode': 200, 'headers': cors_headers, 'body': ''}

    json_headers = {**cors_headers, 'Content-Type': 'application/json'}

    admin_key = os.environ.get('ADMIN_KEY', '')
    headers = event.get('headers', {})
    provided = headers.get('X-Admin-Key') or headers.get('x-admin-key', '')
    if admin_key and provided != admin_key:
        return {
            'statusCode': 403,
            'headers': json_headers,
            'body': json.dumps({'error': 'Доступ запрещён'}, ensure_ascii=False),
            'isBase64Encoded': False
        }

    dsn = os.environ.get('DATABASE_URL')

    if method == 'GET':
        return {
            'statusCode': 200,
            'headers': json_headers,
            'body': json.dumps({'backups': list_backups(dsn)}, ensure_ascii=False),
            'isBase64Encoded': False
        }

    if method != 'POST':
        return {
            'statusCode': 405,
            'headers': json_headers,
            'body': json.dumps({'error': 'Метод не поддерживается'}, ensure_ascii=False),
            'isBase64Encoded': False
        }

    leads, docs = fetch_data(dsn)
    if not leads:
        return {
            'statusCode': 400,
            'headers': json_headers,
            'body': json.dumps({'error': 'В базе пока нет заявок для выгрузки'}, ensure_ascii=False),
            'isBase64Encoded': False
        }

    xlsx_bytes, total_sum = build_workbook(leads, docs)

    now = datetime.now()
    period = now.strftime('%Y-%m')

    s3 = boto3.client(
        's3',
        endpoint_url='https://bucket.poehali.dev',
        aws_access_key_id=os.environ['AWS_ACCESS_KEY_ID'],
        aws_secret_access_key=os.environ['AWS_SECRET_ACCESS_KEY']
    )
    file_key = f"backups/baza_{period}_{int(time.time())}.xlsx"
    s3.put_object(
        Bucket='files',
        Key=file_key,
        Body=xlsx_bytes,
        ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    cdn_url = f"https://cdn.poehali.dev/projects/{os.environ['AWS_ACCESS_KEY_ID']}/bucket/{file_key}"

    conn = psycopg2.connect(dsn)
    cur = conn.cursor()
    cur.execute(
        f"INSERT INTO backups (period, file_url, leads_count, docs_count, total_sum) "
        f"VALUES ('{period}', '{cdn_url}', {len(leads)}, {len(docs)}, {int(total_sum)}) RETURNING id"
    )
    backup_id = cur.fetchone()[0]
    conn.commit()
    cur.close()
    conn.close()

    return {
        'statusCode': 200,
        'headers': json_headers,
        'body': json.dumps({
            'success': True,
            'id': backup_id,
            'url': cdn_url,
            'period': period,
            'leadsCount': len(leads),
            'docsCount': len(docs),
            'totalSum': int(total_sum),
        }, ensure_ascii=False),
        'isBase64Encoded': False
    }
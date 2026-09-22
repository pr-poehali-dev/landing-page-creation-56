import json
import os
import io
import time
import calendar
import base64
from datetime import date
import psycopg2
import boto3
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

MONTHS_RU = ["", "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
             "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]

WEEKDAYS = ["пн", "вт", "ср", "чт", "пт", "сб", "вс"]

PAY_LABELS = {'paid': 'опл', 'unpaid': 'не опл', 'barter': 'бартер',
              'social': 'социальная реклама', 'promo': 'промо реклама', 'unknown': ''}
VIDEO_LABELS = {'ready': 'Присутствует', 'production': 'В процессе изготовления', 'unknown': ''}

HEAD_FILL = PatternFill('solid', fgColor='1E293B')
HEAD_FONT = Font(bold=True, color='FFFFFF', size=10)
TOT_FILL = PatternFill('solid', fgColor='E2E8F0')
WARN_FILL = PatternFill('solid', fgColor='FEE2E2')
THIN = Side(style='thin', color='CBD5E1')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_sheet(ws, year, month, rows, capacity):
    days = calendar.monthrange(year, month)[1]
    ws.cell(row=1, column=1, value=f"{MONTHS_RU[month]} {year}").font = Font(bold=True, size=12)
    for d in range(days):
        wd = date(year, month, d + 1).weekday()
        c = ws.cell(row=1, column=12 + d, value=WEEKDAYS[wd])
        c.alignment = Alignment(horizontal='center')
        c.font = Font(size=8, color='64748B')

    head = ['№', 'Название ролика/Бренд', 'Юрлицо/Физлицо', 'Оплата', 'Хр-ж ролика, сек.',
            'Статус ролика', 'Даты размещения', 'Стоимость за весь период размещения/руб.',
            'Скидка', 'Кол-во дней размещения в месяце',
            'Стоимость  за текущий период размещения/ руб.']
    for i, title in enumerate(head, 1):
        c = ws.cell(row=2, column=i, value=title)
        c.fill = HEAD_FILL; c.font = HEAD_FONT
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for d in range(days):
        c = ws.cell(row=2, column=12 + d, value=d + 1)
        c.fill = HEAD_FILL; c.font = HEAD_FONT
        c.alignment = Alignment(horizontal='center')
    ws.row_dimensions[2].height = 42

    load = [0] * days
    r = 3
    for idx, row in enumerate(rows, 1):
        (row_no, brand, legal, agency, pay, vid, dur, period,
         days_count, price_total, discount, amount, sec_str) = row
        ws.cell(row=r, column=1, value=row_no or idx)
        ws.cell(row=r, column=2, value=brand)
        ws.cell(row=r, column=3, value=agency or legal or '')
        ws.cell(row=r, column=4, value=PAY_LABELS.get(pay, ''))
        ws.cell(row=r, column=5, value=dur or None)
        ws.cell(row=r, column=6, value=VIDEO_LABELS.get(vid, ''))
        ws.cell(row=r, column=7, value=period or '')
        ws.cell(row=r, column=8, value=price_total or None)
        ws.cell(row=r, column=9, value=float(discount or 0))
        ws.cell(row=r, column=10, value=days_count or None)
        ws.cell(row=r, column=11, value=amount or None)

        secs = []
        if sec_str:
            for part in str(sec_str).split(','):
                try:
                    secs.append(int(part))
                except ValueError:
                    secs.append(0)
        for d in range(days):
            v = secs[d] if d < len(secs) else 0
            if v:
                load[d] += v
                cell = ws.cell(row=r, column=12 + d, value=v)
                cell.alignment = Alignment(horizontal='center')
                cell.font = Font(size=9)
        r += 1

    def total_row(label, values, fill, color=None):
        nonlocal r
        c = ws.cell(row=r, column=1, value=label)
        c.font = Font(bold=True, size=10)
        for col in range(1, 12 + days):
            ws.cell(row=r, column=col).fill = fill
        for d, v in enumerate(values):
            cell = ws.cell(row=r, column=12 + d, value=v)
            cell.alignment = Alignment(horizontal='center')
            cell.font = Font(bold=True, size=9, color=color or '000000')
        r += 1

    total_amount = sum((x[11] or 0) for x in rows)
    ws.cell(row=r, column=11, value=total_amount).font = Font(bold=True)
    total_row('Итого загрузка, сек', load, TOT_FILL)
    total_row('Остаток времени, сек', [max(capacity - x, 0) for x in load], TOT_FILL, '047857')
    pct = [round(x / capacity * 100) if capacity else 0 for x in load]
    total_row('% заполнения блока', pct, TOT_FILL)
    for d, p in enumerate(pct):
        if p >= 90:
            ws.cell(row=r - 1, column=12 + d).fill = WARN_FILL

    widths = {'A': 5, 'B': 34, 'C': 24, 'D': 16, 'E': 11, 'F': 20, 'G': 22,
              'H': 16, 'I': 9, 'J': 14, 'K': 18}
    for col, wdt in widths.items():
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = 'L3'

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=11 + days):
        for cell in row:
            cell.border = BORDER
    return total_amount


def log_export(event, what):
    """Пишет выгрузку данных в журнал действий"""
    headers = event.get('headers') or {}
    token = headers.get('X-Session-Token') or headers.get('x-session-token')
    req = event.get('requestContext') or {}
    ident = req.get('identity') or {}
    fwd = headers.get('X-Forwarded-For') or headers.get('x-forwarded-for') or ''
    ip = ((fwd.split(',')[0].strip() if fwd else '') or ident.get('sourceIp') or '')[:45]
    c = psycopg2.connect(os.environ['DATABASE_URL'])
    k = c.cursor()
    sid, name = 'NULL', 'NULL'
    if token:
        safe = str(token).replace("'", "''")
        k.execute(
            "SELECT st.id, st.name FROM staff_sessions s JOIN staff st ON st.id = s.staff_id "
            f"WHERE s.token = '{safe}' AND s.revoked = FALSE"
        )
        r = k.fetchone()
        if r:
            sid = str(r[0])
            name = "'" + str(r[1]).replace("'", "''") + "'"
    text = str(what)[:500].replace("'", "''")
    k.execute(
        "INSERT INTO audit_log (staff_id, staff_name, action, entity, details, ip, severity) "
        f"VALUES ({sid}, {name}, 'export', 'system', '{text}', '{ip}', 'warning')"
    )
    c.commit()
    k.close()
    c.close()


def session_ok(event):
    """Доступ по мастер-паролю или активной сессии сотрудника"""
    master = os.environ.get('ADMIN_KEY', '')
    headers = event.get('headers') or {}
    provided = headers.get('X-Admin-Key') or headers.get('x-admin-key', '')
    if master and provided == master:
        return True
    token = headers.get('X-Session-Token') or headers.get('x-session-token')
    if not token:
        return False
    safe = str(token).replace("'", "''")
    c = psycopg2.connect(os.environ['DATABASE_URL'])
    k = c.cursor()
    k.execute(
        "SELECT 1 FROM staff_sessions s JOIN staff st ON st.id = s.staff_id "
        f"WHERE s.token = '{safe}' AND s.revoked = FALSE "
        "AND s.expires_at > CURRENT_TIMESTAMP AND st.active = TRUE"
    )
    ok = k.fetchone() is not None
    k.close()
    c.close()
    return ok


def handler(event: dict, context) -> dict:
    """Выгружает медиаплан размещений в Excel в привычном формате: лист на каждый месяц с сеткой дней и загрузкой экрана"""
    method = event.get('httpMethod', 'POST')
    cors = {'Access-Control-Allow-Origin': '*',
            'Access-Control-Allow-Methods': 'POST, OPTIONS',
            'Access-Control-Allow-Headers': 'Content-Type, X-Admin-Key'}

    if method == 'OPTIONS':
        return {'statusCode': 200, 'headers': cors, 'body': ''}
    if not session_ok(event):
        return {'statusCode': 403, 'headers': {**cors, 'Content-Type': 'application/json'},
                'body': json.dumps({'error': 'Доступ запрещён'}, ensure_ascii=False)}

    log_export(event, 'Выгрузка медиаплана в Excel')
    body = json.loads(event.get('body', '{}'))
    only_year = body.get('year')
    only_month = body.get('month')

    conn = psycopg2.connect(os.environ['DATABASE_URL'])
    cur = conn.cursor()
    cur.execute("SELECT daily_capacity_sec FROM screen_settings WHERE id = 1")
    st = cur.fetchone()
    capacity = st[0] if st else 300

    if only_year and only_month:
        cur.execute(
            "SELECT DISTINCT plan_year, plan_month FROM placements "
            f"WHERE plan_year = {int(only_year)} AND plan_month = {int(only_month)}"
        )
    else:
        cur.execute("SELECT DISTINCT plan_year, plan_month FROM placements ORDER BY plan_year, plan_month")
    periods = cur.fetchall()

    if not periods:
        cur.close(); conn.close()
        return {'statusCode': 400, 'headers': {**cors, 'Content-Type': 'application/json'},
                'body': json.dumps({'error': 'Нет данных для выгрузки'}, ensure_ascii=False)}

    wb = Workbook()
    wb.remove(wb.active)
    summary = []

    for (y, m) in periods:
        cur.execute(
            "SELECT row_no, brand, legal_entity, agency, payment_type, video_status, "
            "duration_sec, period_text, days_count, price_total, discount, amount_month, day_seconds "
            f"FROM placements WHERE plan_year = {y} AND plan_month = {m} "
            "ORDER BY COALESCE(row_no, 999), id"
        )
        rows = cur.fetchall()
        ws = wb.create_sheet(f"{MONTHS_RU[m]} {y}"[:31])
        total = build_sheet(ws, y, m, rows, capacity)
        summary.append((y, m, len(rows), total))

    cur.execute("SELECT fact_year, fact_month, amount FROM revenue_facts ORDER BY fact_year, fact_month")
    facts = cur.fetchall()
    cur.close()
    conn.close()

    ws = wb.create_sheet('Сводная', 0)
    ws.append(['Период', 'Роликов', 'Сумма плана, ₽', 'Факт выручки, ₽'])
    for i in range(1, 5):
        c = ws.cell(row=1, column=i)
        c.fill = HEAD_FILL; c.font = HEAD_FONT
        c.alignment = Alignment(horizontal='center')
    fact_map = {(f[0], f[1]): float(f[2]) for f in facts}
    for y, m, cnt, total in summary:
        ws.append([f"{MONTHS_RU[m]} {y}", cnt, total, fact_map.get((y, m), '')])
    r = ws.max_row + 1
    ws.cell(row=r, column=1, value='ИТОГО').font = Font(bold=True)
    ws.cell(row=r, column=2, value=sum(s[2] for s in summary)).font = Font(bold=True)
    ws.cell(row=r, column=3, value=sum(s[3] for s in summary)).font = Font(bold=True)
    for i in range(1, 5):
        ws.cell(row=r, column=i).fill = TOT_FILL

    if facts:
        ws.append([])
        ws.append(['Факт выручки прошлых лет'])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=11)
        ws.append(['Период', 'Сумма, ₽'])
        for fy, fm, amount in facts:
            ws.append([f"{MONTHS_RU[fm]} {fy}", float(amount)])

    for col, wdt in {'A': 20, 'B': 12, 'C': 18, 'D': 18}.items():
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = 'A2'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    data = buf.read()

    s3 = boto3.client('s3', endpoint_url='https://bucket.poehali.dev',
                      aws_access_key_id=os.environ['AWS_ACCESS_KEY_ID'],
                      aws_secret_access_key=os.environ['AWS_SECRET_ACCESS_KEY'])
    key = f"mediaplan/mediaplan_{int(time.time())}.xlsx"
    s3.put_object(Bucket='files', Key=key, Body=data,
                  ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    url = f"https://cdn.poehali.dev/projects/{os.environ['AWS_ACCESS_KEY_ID']}/bucket/{key}"

    return {'statusCode': 200, 'headers': {**cors, 'Content-Type': 'application/json'},
            'body': json.dumps({'url': url, 'sheets': len(summary)}, ensure_ascii=False),
            'isBase64Encoded': False}
